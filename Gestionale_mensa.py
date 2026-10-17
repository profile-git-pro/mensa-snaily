import streamlit as st
import json
import os
import io
import pandas as pd
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.drawing.image import Image as OpenpyxlImage
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from reportlab.lib.pagesizes import A4, landscape
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, Image as ReportLabImage
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib import colors
from PIL import Image, ImageDraw, ImageFont

# ==========================================
# CONFIGURAZIONE PAGINA E TEMA CUSTO
# ==========================================
st.set_page_config(
    page_title="Scuole Snaily - Gestionale Menù",
    page_icon="🍲",
    layout="wide",
    initial_sidebar_state="expanded"
)

# Custom CSS basato sui colori richiesti (#024b34 sfondo, #C5A067 oro per i caratteri/accenti)
CUSTOM_CSS = """
<style>
    /* Importazione Font Eleganti */
    @import url('https://fonts.googleapis.com/css2?family=Cinzel:wght@600;700&family=Montserrat:wght@400;500;600;700&display=swap');

    /* Sfondo Principale e Colore Testo Base */
    .stApp, [data-testid="stAppViewContainer"] {
        background-color: #024b34 !important;
        color: #f1e6d0 !important;
        font-family: 'Montserrat', sans-serif !important;
    }

    /* Header e Barra Superiore */
    [data-testid="stHeader"] {
        background-color: rgba(2, 75, 52, 0.95) !important;
    }

    /* Sidebar Personale */
    [data-testid="stSidebar"] {
        background-color: #013524 !important;
        border-right: 2px solid #C5A067 !important;
    }
    [data-testid="stSidebar"] p, [data-testid="stSidebar"] span, [data-testid="stSidebar"] label {
        color: #f1e6d0 !important;
    }

    /* ==========================================
       TITOLI ED INTESTAZIONI UNIVERSALI (#C5A067 ORO)
       Garantisce il colore oro #C5A067 su qualsiasi livello di intestazione
       ========================================== */
    h1, h2, h3, h4, h5, h6,
    h1 *, h2 *, h3 *, h4 *, h5 *, h6 *,
    .stMarkdown h1, .stMarkdown h2, .stMarkdown h3, .stMarkdown h4, .stMarkdown h5, .stMarkdown h6,
    .stMarkdown h1 *, .stMarkdown h2 *, .stMarkdown h3 *, .stMarkdown h4 *, .stMarkdown h5 *, .stMarkdown h6 *,
    [data-testid="stMarkdownContainer"] h1,
    [data-testid="stMarkdownContainer"] h2,
    [data-testid="stMarkdownContainer"] h3,
    [data-testid="stMarkdownContainer"] h4,
    [data-testid="stMarkdownContainer"] h5,
    [data-testid="stMarkdownContainer"] h6,
    [data-testid="stMarkdownContainer"] h1 *,
    [data-testid="stMarkdownContainer"] h2 *,
    [data-testid="stMarkdownContainer"] h3 *,
    [data-testid="stMarkdownContainer"] h4 *,
    [data-testid="stMarkdownContainer"] h5 *,
    [data-testid="stMarkdownContainer"] h6 *,
    [data-testid="stHeading"],
    [data-testid="stHeading"] *,
    [data-testid="stHeader"],
    [data-testid="stHeader"] *,
    .stHeader, .stSubheader, .stTitle,
    .stHeader *, .stSubheader *, .stTitle *,
    .hero-title, .info-card h2, .info-card h3, .info-card h4 {
        color: #C5A067 !important;
        font-family: 'Cinzel', Georgia, serif !important;
        font-weight: 700 !important;
        letter-spacing: 0.5px !important;
    }

    /* Testo Paragrafi, Dettagli e Didascalie (Crema) */
    p, .stCaption, [data-testid="stMarkdownContainer"] p {
        color: #e8dcc4 !important;
    }

    /* Hero Banner Stile Template Nicepage */
    .hero-card {
        background: linear-gradient(135deg, rgba(13, 92, 67, 0.95) 0%, rgba(1, 53, 36, 0.95) 100%);
        border: 2px solid #C5A067;
        border-radius: 16px;
        padding: 24px 30px;
        margin-bottom: 25px;
        box-shadow: 0 10px 25px rgba(0,0,0,0.4);
        display: flex;
        align-items: center;
        gap: 25px;
    }
    .hero-title {
        color: #C5A067 !important;
        font-family: 'Cinzel', serif !important;
        font-size: 28px !important;
        font-weight: 700;
        margin: 0 0 8px 0 !important;
    }
    .hero-subtitle {
        color: #ffffff !important;
        font-size: 15px !important;
        margin: 0 !important;
    }
    .badge-gold {
        background-color: #C5A067;
        color: #024b34;
        padding: 4px 12px;
        border-radius: 20px;
        font-weight: 700;
        font-size: 12px;
        display: inline-block;
        margin-top: 10px;
    }

    /* Card di Dettaglio / Contenitori */
    .info-card {
        background-color: rgba(10, 84, 59, 0.6);
        border: 1px solid rgba(197, 160, 103, 0.4);
        border-radius: 12px;
        padding: 18px;
        margin-bottom: 15px;
        box-shadow: 0 4px 12px rgba(0,0,0,0.2);
    }
    .info-card h2, .info-card h3, .info-card h4 {
        color: #C5A067 !important;
    }

    /* Pulsanti */
    .stButton > button, div[data-testid="stFormSubmitButton"] > button, .stDownloadButton > button {
        background: linear-gradient(135deg, #C5A067 0%, #d4b27d 100%) !important;
        color: #024b34 !important;
        font-weight: 700 !important;
        font-family: 'Montserrat', sans-serif !important;
        border-radius: 8px !important;
        border: 1px solid #e2be87 !important;
        padding: 10px 22px !important;
        box-shadow: 0 4px 12px rgba(0,0,0,0.3) !important;
        transition: all 0.3s ease !important;
    }
    .stButton > button:hover, div[data-testid="stFormSubmitButton"] > button:hover, .stDownloadButton > button:hover {
        background: linear-gradient(135deg, #e2be87 0%, #ffffff 100%) !important;
        color: #01281b !important;
        border-color: #ffffff !important;
        transform: translateY(-2px) !important;
        box-shadow: 0 6px 18px rgba(197, 160, 103, 0.5) !important;
    }

    /* ==========================================
       CASELLE DI SCELTA, INPUT, DROPDOWN E POPOVER
       Testo nero (#000000) su sfondo bianco (#ffffff) per massimo contrasto
       ========================================== */

    /* Etichette (Labels) sopra i campi - In Oro #C5A067 */
    label, .stSelectbox label, .stTextInput label, .stNumberInput label, .stCheckbox label, .stRadio label {
        color: #C5A067 !important;
        font-weight: 600 !important;
        font-size: 14px !important;
    }

    /* Contenitori Input e Selectbox (Caselle di Scelta) */
    div[data-baseweb="select"] > div,
    .stTextInput > div > div > input,
    .stNumberInput > div > div > input,
    div[data-baseweb="input"] > input,
    div[data-baseweb="base-input"] > input,
    .stSelectbox div[role="combobox"] {
        background-color: #ffffff !important;
        color: #000000 !important;
        font-weight: 600 !important;
        font-size: 14px !important;
        border: 2px solid #C5A067 !important;
        border-radius: 8px !important;
    }

    /* Testo all'interno di tutti i selectbox, input e opzioni */
    div[data-baseweb="select"] *,
    div[data-baseweb="select"] span,
    div[data-baseweb="select"] div,
    .stSelectbox *,
    .stTextInput input,
    .stNumberInput input {
        color: #000000 !important;
        font-weight: 600 !important;
    }

    /* Menu a tendina aperto (Popover / Dropdown list) */
    div[data-baseweb="popover"],
    div[data-baseweb="popover"] *,
    div[data-baseweb="menu"],
    div[data-baseweb="menu"] *,
    ul[role="listbox"],
    ul[role="listbox"] *,
    li[role="option"],
    li[role="option"] *,
    div[role="option"],
    div[role="option"] * {
        background-color: #ffffff !important;
        color: #000000 !important;
        font-weight: 600 !important;
    }

    /* Hover e selezione attiva nelle opzioni del menu a tendina */
    li[role="option"]:hover,
    div[role="option"]:hover,
    [data-baseweb="menu"] li:hover,
    [data-baseweb="menu"] div:hover {
        background-color: #e8f2ee !important;
        color: #024b34 !important;
        font-weight: 700 !important;
    }

    /* Frecce e Icone dentro i Selectbox */
    div[data-baseweb="select"] svg {
        fill: #024b34 !important;
    }

    /* Checkbox e Radio Label Color */
    .stCheckbox span, .stRadio span {
        color: #f1e6d0 !important;
    }

    /* Tab Layout */
    .stTabs [data-baseweb="tab-list"] {
        gap: 8px;
        background-color: #013524;
        padding: 8px;
        border-radius: 10px;
        border: 1px solid rgba(197, 160, 103, 0.3);
    }
    .stTabs [data-baseweb="tab"] {
        background-color: #08583e !important;
        color: #C5A067 !important;
        border-radius: 6px !important;
        padding: 8px 18px !important;
        font-weight: 600 !important;
        border: 1px solid transparent !important;
    }
    .stTabs [aria-selected="true"] {
        background-color: #C5A067 !important;
        color: #024b34 !important;
        font-weight: 700 !important;
    }

    /* Tabelle e Dataframe */
    [data-testid="stDataFrame"] {
        border: 1px solid #C5A067 !important;
        border-radius: 10px !important;
        overflow: hidden !important;
    }
    
    /* Divider e Separatori */
    hr {
        border-color: rgba(197, 160, 103, 0.4) !important;
    }
</style>
"""

st.markdown(CUSTOM_CSS, unsafe_allow_html=True)

DB_FILE = "mensa_db.json"
LOGO_PATH = "Logo_Snaily.png"

# Verifichiamo la presenza del logo Snaily
def ensure_default_logo():
    if not os.path.exists(LOGO_PATH):
        alt_paths = ["/workspace/artifacts/Logo_Snaily.png", "/workspace/knowledge/Logo_Snaily.png"]
        found = False
        for alt in alt_paths:
            if os.path.exists(alt):
                import shutil
                shutil.copy(alt, LOGO_PATH)
                found = True
                break
        if not found:
            try:
                width, height = 126, 167
                img = Image.new('RGBA', (width, height), (2, 75, 52, 255))
                draw = ImageDraw.Draw(img)
                draw.rectangle([5, 5, width-5, height-5], outline=(197, 160, 103, 255), width=3)
                img.save(LOGO_PATH)
            except Exception:
                pass

ensure_default_logo()

# ==========================================
# INIZIALIZZAZIONE DATABASE LOCALE JSON
# ==========================================
DEFAULT_ANAGRAFICA = [
    {"id": 1, "categoria": "Primo", "nome": "Pasta al pomodoro biologico", "allergeni": "Glutine", "nido": 40, "infanzia": 50, "primaria": 60},
    {"id": 2, "categoria": "Primo", "nome": "Risotto alle verdure di stagione", "allergeni": "Nessuno", "nido": 45, "infanzia": 55, "primaria": 65},
    {"id": 3, "categoria": "Primo", "nome": "Pasta e fagioli", "allergeni": "Glutine", "nido": 45, "infanzia": 60, "primaria": 75},
    {"id": 4, "categoria": "Primo", "nome": "Pasta e lenticchie", "allergeni": "Glutine", "nido": 45, "infanzia": 60, "primaria": 75},
    {"id": 5, "categoria": "Primo", "nome": "Pasta e ceci", "allergeni": "Glutine", "nido": 45, "infanzia": 60, "primaria": 75},
    {"id": 6, "categoria": "Primo", "nome": "Passato di verdure con pasta", "allergeni": "Glutine", "nido": 40, "infanzia": 50, "primaria": 65},
    {"id": 7, "categoria": "Primo", "nome": "Minestrone di verdure con riso", "allergeni": "Nessuno", "nido": 45, "infanzia": 55, "primaria": 70},
    {"id": 8, "categoria": "Primo", "nome": "Pasta al ragù di carne Podolica", "allergeni": "Glutine, Sedano", "nido": 45, "infanzia": 55, "primaria": 70},
    {"id": 9, "categoria": "Primo", "nome": "Gnocchi biologici al pomodoro", "allergeni": "Glutine", "nido": 50, "infanzia": 70, "primaria": 90},
    {"id": 10, "categoria": "Secondo", "nome": "Polpette al sugo di carne Podolica", "allergeni": "Glutine, Uova", "nido": 60, "infanzia": 70, "primaria": 90},
    {"id": 11, "categoria": "Secondo", "nome": "Petto di pollo biologico al limone", "allergeni": "Nessuno", "nido": 60, "infanzia": 80, "primaria": 100},
    {"id": 12, "categoria": "Secondo", "nome": "Spezzatino di vitellone Podolico con patate", "allergeni": "Sedano", "nido": 60, "infanzia": 80, "primaria": 100},
    {"id": 13, "categoria": "Secondo", "nome": "Filetto di platessa al forno", "allergeni": "Pesce", "nido": 70, "infanzia": 90, "primaria": 110},
    {"id": 14, "categoria": "Secondo", "nome": "Bastoncini di merluzzo al forno", "allergeni": "Glutine, Pesce", "nido": 70, "infanzia": 80, "primaria": 100},
    {"id": 15, "categoria": "Secondo", "nome": "Frittata al forno con spinaci", "allergeni": "Uova", "nido": 50, "infanzia": 65, "primaria": 80},
    {"id": 16, "categoria": "Secondo", "nome": "Bocconcini di mozzarella fresca DOP", "allergeni": "Latte", "nido": 40, "infanzia": 60, "primaria": 80},
    {"id": 17, "categoria": "Contorno", "nome": "Patate al forno con rosmarino", "allergeni": "Nessuno", "nido": 70, "infanzia": 90, "primaria": 110},
    {"id": 18, "categoria": "Contorno", "nome": "Purè di patate fresche", "allergeni": "Latte", "nido": 80, "infanzia": 100, "primaria": 120},
    {"id": 19, "categoria": "Contorno", "nome": "Carote julienne all'olio EVO", "allergeni": "Nessuno", "nido": 60, "infanzia": 80, "primaria": 90},
    {"id": 20, "categoria": "Contorno", "nome": "Spinaci ripassati all'olio EVO", "allergeni": "Nessuno", "nido": 60, "infanzia": 80, "primaria": 90},
    {"id": 21, "categoria": "Contorno", "nome": "Pomodori freschi all'insalata", "allergeni": "Nessuno", "nido": 60, "infanzia": 80, "primaria": 90},
    {"id": 22, "categoria": "Pane/Frutta", "nome": "Pane tipo '1' artigianale + Mela biologica", "allergeni": "Glutine", "nido": 120, "infanzia": 150, "primaria": 180},
    {"id": 23, "categoria": "Pane/Frutta", "nome": "Pane tipo '1' artigianale + Pera di stagione", "allergeni": "Glutine", "nido": 120, "infanzia": 150, "primaria": 180},
    {"id": 24, "categoria": "Pane/Frutta", "nome": "Pane tipo '1' artigianale + Clementine di Calabria IGP", "allergeni": "Glutine", "nido": 120, "infanzia": 150, "primaria": 180},
    {"id": 25, "categoria": "Pane/Frutta", "nome": "Pane tipo '1' artigianale + Arancia biologica", "allergeni": "Glutine", "nido": 120, "infanzia": 150, "primaria": 180},
]

DEFAULT_MENU = {
    "Invernale": {
        "Settimana 1": {
            "Lunedì": {"Primo": "Pasta al pomodoro biologico", "Secondo": "Polpette al sugo di carne Podolica", "Contorno": "Carote julienne all'olio EVO", "Pane/Frutta": "Pane tipo '1' artigianale + Mela biologica"},
            "Martedì": {"Primo": "Risotto alle verdure di stagione", "Secondo": "Filetto di platessa al forno", "Contorno": "Patate al forno con rosmarino", "Pane/Frutta": "Pane tipo '1' artigianale + Pera di stagione"},
            "Mercoledì": {"Primo": "Pasta e fagioli", "Secondo": "Bocconcini di mozzarella fresca DOP", "Contorno": "Spinaci ripassati all'olio EVO", "Pane/Frutta": "Pane tipo '1' artigianale + Clementine di Calabria IGP"},
            "Giovedì": {"Primo": "Pasta e lenticchie", "Secondo": "Frittata al forno con spinaci", "Contorno": "Pomodori freschi all'insalata", "Pane/Frutta": "Pane tipo '1' artigianale + Arancia biologica"},
            "Venerdì": {"Primo": "Minestrone di verdure con riso", "Secondo": "Bastoncini di merluzzo al forno", "Contorno": "Purè di patate fresche", "Pane/Frutta": "Pane tipo '1' artigianale + Mela biologica"}
        },
        "Settimana 2": {
            "Lunedì": {"Primo": "Pasta e ceci", "Secondo": "Petto di pollo biologico al limone", "Contorno": "Purè di patate fresche", "Pane/Frutta": "Pane tipo '1' artigianale + Clementine di Calabria IGP"},
            "Martedì": {"Primo": "Pasta al ragù di carne Podolica", "Secondo": "Frittata al forno con spinaci", "Contorno": "Carote julienne all'olio EVO", "Pane/Frutta": "Pane tipo '1' artigianale + Mela biologica"},
            "Mercoledì": {"Primo": "Passato di verdure con pasta", "Secondo": "Filetto di platessa al forno", "Contorno": "Patate al forno con rosmarino", "Pane/Frutta": "Pane tipo '1' artigianale + Pera di stagione"},
            "Giovedì": {"Primo": "Gnocchi biologici al pomodoro", "Secondo": "Bocconcini di mozzarella fresca DOP", "Contorno": "Spinaci ripassati all'olio EVO", "Pane/Frutta": "Pane tipo '1' artigianale + Arancia biologica"},
            "Venerdì": {"Primo": "Risotto alle verdure di stagione", "Secondo": "Bastoncini di merluzzo al forno", "Contorno": "Pomodori freschi all'insalata", "Pane/Frutta": "Pane tipo '1' artigianale + Mela biologica"}
        },
        "Settimana 3": {
            "Lunedì": {"Primo": "Pasta e fagioli", "Secondo": "Spezzatino di vitellone Podolico con patate", "Contorno": "Carote julienne all'olio EVO", "Pane/Frutta": "Pane tipo '1' artigianale + Pera di stagione"},
            "Martedì": {"Primo": "Pasta al pomodoro biologico", "Secondo": "Bastoncini di merluzzo al forno", "Contorno": "Purè di patate fresche", "Pane/Frutta": "Pane tipo '1' artigianale + Clementine di Calabria IGP"},
            "Mercoledì": {"Primo": "Minestrone di verdure con riso", "Secondo": "Polpette al sugo di carne Podolica", "Contorno": "Spinaci ripassati all'olio EVO", "Pane/Frutta": "Pane tipo '1' artigianale + Mela biologica"},
            "Giovedì": {"Primo": "Pasta e lenticchie", "Secondo": "Bocconcini di mozzarella fresca DOP", "Contorno": "Patate al forno con rosmarino", "Pane/Frutta": "Pane tipo '1' artigianale + Arancia biologica"},
            "Venerdì": {"Primo": "Pasta al ragù di carne Podolica", "Secondo": "Filetto di platessa al forno", "Contorno": "Pomodori freschi all'insalata", "Pane/Frutta": "Pane tipo '1' artigianale + Pera di stagione"}
        },
        "Settimana 4": {
            "Lunedì": {"Primo": "Risotto alle verdure di stagione", "Secondo": "Frittata al forno con spinaci", "Contorno": "Patate al forno con rosmarino", "Pane/Frutta": "Pane tipo '1' artigianale + Clementine di Calabria IGP"},
            "Martedì": {"Primo": "Pasta e ceci", "Secondo": "Petto di pollo biologico al limone", "Contorno": "Spinaci ripassati all'olio EVO", "Pane/Frutta": "Pane tipo '1' artigianale + Mela biologica"},
            "Mercoledì": {"Primo": "Passato di verdure con pasta", "Secondo": "Bastoncini di merluzzo al forno", "Contorno": "Carote julienne all'olio EVO", "Pane/Frutta": "Pane tipo '1' artigianale + Pera di stagione"},
            "Giovedì": {"Primo": "Gnocchi biologici al pomodoro", "Secondo": "Polpette al sugo di carne Podolica", "Contorno": "Purè di patate fresche", "Pane/Frutta": "Pane tipo '1' artigianale + Arancia biologica"},
            "Venerdì": {"Primo": "Pasta e fagioli", "Secondo": "Filetto di platessa al forno", "Contorno": "Pomodori freschi all'insalata", "Pane/Frutta": "Pane tipo '1' artigianale + Clementine di Calabria IGP"}
        }
    },
    "Estivo": {
        "Settimana 1": {
            "Lunedì": {"Primo": "Pasta al pomodoro biologico", "Secondo": "Petto di pollo biologico al limone", "Contorno": "Pomodori freschi all'insalata", "Pane/Frutta": "Pane tipo '1' artigianale + Mela biologica"},
            "Martedì": {"Primo": "Risotto alle verdure di stagione", "Secondo": "Filetto di platessa al forno", "Contorno": "Carote julienne all'olio EVO", "Pane/Frutta": "Pane tipo '1' artigianale + Pera di stagione"},
            "Mercoledì": {"Primo": "Pasta e ceci", "Secondo": "Bocconcini di mozzarella fresca DOP", "Contorno": "Patate al forno con rosmarino", "Pane/Frutta": "Pane tipo '1' artigianale + Clementine di Calabria IGP"},
            "Giovedì": {"Primo": "Gnocchi biologici al pomodoro", "Secondo": "Frittata al forno con spinaci", "Contorno": "Pomodori freschi all'insalata", "Pane/Frutta": "Pane tipo '1' artigianale + Arancia biologica"},
            "Venerdì": {"Primo": "Minestrone di verdure con riso", "Secondo": "Bastoncini di merluzzo al forno", "Contorno": "Carote julienne all'olio EVO", "Pane/Frutta": "Pane tipo '1' artigianale + Mela biologica"}
        },
        "Settimana 2": {
            "Lunedì": {"Primo": "Pasta e lenticchie", "Secondo": "Polpette al sugo di carne Podolica", "Contorno": "Carote julienne all'olio EVO", "Pane/Frutta": "Pane tipo '1' artigianale + Pera di stagione"},
            "Martedì": {"Primo": "Risotto alle verdure di stagione", "Secondo": "Filetto di platessa al forno", "Contorno": "Pomodori freschi all'insalata", "Pane/Frutta": "Pane tipo '1' artigianale + Clementine di Calabria IGP"},
            "Mercoledì": {"Primo": "Pasta al pomodoro biologico", "Secondo": "Petto di pollo biologico al limone", "Contorno": "Patate al forno con rosmarino", "Pane/Frutta": "Pane tipo '1' artigianale + Mela biologica"},
            "Giovedì": {"Primo": "Passato di verdure con pasta", "Secondo": "Bocconcini di mozzarella fresca DOP", "Contorno": "Carote julienne all'olio EVO", "Pane/Frutta": "Pane tipo '1' artigianale + Arancia biologica"},
            "Venerdì": {"Primo": "Pasta e fagioli", "Secondo": "Bastoncini di merluzzo al forno", "Contorno": "Purè di patate fresche", "Pane/Frutta": "Pane tipo '1' artigianale + Pera di stagione"}
        },
        "Settimana 3": {
            "Lunedì": {"Primo": "Gnocchi biologici al pomodoro", "Secondo": "Frittata al forno con spinaci", "Contorno": "Pomodori freschi all'insalata", "Pane/Frutta": "Pane tipo '1' artigianale + Clementine di Calabria IGP"},
            "Martedì": {"Primo": "Pasta e ceci", "Secondo": "Filetto di platessa al forno", "Contorno": "Carote julienne all'olio EVO", "Pane/Frutta": "Pane tipo '1' artigianale + Mela biologica"},
            "Mercoledì": {"Primo": "Risotto alle verdure di stagione", "Secondo": "Polpette al sugo di carne Podolica", "Contorno": "Purè di patate fresche", "Pane/Frutta": "Pane tipo '1' artigianale + Pera di stagione"},
            "Giovedì": {"Primo": "Pasta al ragù di carne Podolica", "Secondo": "Bocconcini di mozzarella fresca DOP", "Contorno": "Patate al forno con rosmarino", "Pane/Frutta": "Pane tipo '1' artigianale + Arancia biologica"},
            "Venerdì": {"Primo": "Minestrone di verdure con riso", "Secondo": "Bastoncini di merluzzo al forno", "Contorno": "Pomodori freschi all'insalata", "Pane/Frutta": "Pane tipo '1' artigianale + Clementine di Calabria IGP"}
        },
        "Settimana 4": {
            "Lunedì": {"Primo": "Pasta al pomodoro biologico", "Secondo": "Petto di pollo biologico al limone", "Contorno": "Patate al forno con rosmarino", "Pane/Frutta": "Pane tipo '1' artigianale + Mela biologica"},
            "Martedì": {"Primo": "Pasta e lenticchie", "Secondo": "Filetto di platessa al forno", "Contorno": "Carote julienne all'olio EVO", "Pane/Frutta": "Pane tipo '1' artigianale + Pera di stagione"},
            "Mercoledì": {"Primo": "Passato di verdure con pasta", "Secondo": "Frittata al forno con spinaci", "Contorno": "Pomodori freschi all'insalata", "Pane/Frutta": "Pane tipo '1' artigianale + Clementine di Calabria IGP"},
            "Giovedì": {"Primo": "Risotto alle verdure di stagione", "Secondo": "Spezzatino di vitellone Podolico con patate", "Contorno": "Purè di patate fresche", "Pane/Frutta": "Pane tipo '1' artigianale + Arancia biologica"},
            "Venerdì": {"Primo": "Pasta e fagioli", "Secondo": "Bastoncini di merluzzo al forno", "Contorno": "Carote julienne all'olio EVO", "Pane/Frutta": "Pane tipo '1' artigianale + Mela biologica"}
        }
    }
}

def load_data():
    if os.path.exists(DB_FILE):
        try:
            with open(DB_FILE, "r", encoding="utf-8") as f:
                loaded = json.load(f)
                if "menu" in loaded:
                    for stag in ["Invernale", "Estivo"]:
                        if stag not in loaded["menu"]:
                            loaded["menu"][stag] = DEFAULT_MENU[stag]
                        else:
                            for sett in ["Settimana 1", "Settimana 2", "Settimana 3", "Settimana 4"]:
                                if sett not in loaded["menu"][stag]:
                                    loaded["menu"][stag][sett] = DEFAULT_MENU[stag].get(sett, {})
                return loaded
        except Exception:
            pass
    return {"portate": DEFAULT_ANAGRAFICA, "menu": DEFAULT_MENU}

def save_data(data):
    with open(DB_FILE, "w", encoding="utf-8") as f:
        json.dump(data, f, ensure_ascii=False, indent=4)

data = load_data()

# Mappa allergeni velocemente dall'anagrafica
allergeni_map = {p["nome"]: p.get("allergeni", "Nessuno") for p in data.get("portate", [])}

# ==========================================
# BARRA LATERALE (SIDEBAR) & NAVIGAZIONE
# ==========================================
if os.path.exists(LOGO_PATH):
    st.sidebar.image(LOGO_PATH, width=130)

st.sidebar.markdown("<h2 style='color:#C5A067 !important; margin-bottom:2px;'>Scuole Snaily<br>Servizio Mensa</h2>", unsafe_allow_html=True)
st.sidebar.markdown("<p style='color:#f1e6d0 !important; font-size:13px;'>Piano Nutrizionale</p>", unsafe_allow_html=True)

menu_option = st.sidebar.radio(
    "Navigazione",
    ["📋 Visualizza Menù Settimanale", "✏️ Imposta Menù Settimanale", "🥦 Gestisci Anagrafica Portate"]
)

def render_dataframe(df):
    try:
        st.dataframe(df, width="stretch", hide_index=True)
    except TypeError:
        st.dataframe(df, use_container_width=True, hide_index=True)

# ==========================================
# PAGINA 1: VISUALIZZA E STAMPA MENU
# ==========================================
if menu_option == "📋 Visualizza Menù Settimanale":
    # HERO BANNER STILE NICEPAGE TEMPLATE
    st.markdown(f"""
    <div class="hero-card">
        <div>
            <div class="hero-title">SCUOLE SNAILY</div>
            <div class="hero-subtitle">Gestionale Menù & Piano Nutrizionale Mensa</div>
            <span class="badge-gold">✓ Conforme L.R. Calabria n. 24/2024 & ASP Cosenza</span>
        </div>
    </div>
    """, unsafe_allow_html=True)
    
    col_s1, col_s2, col_s3 = st.columns([2, 2, 3])
    with col_s1:
        stagione_sel = st.selectbox("Seleziona Stagione", ["Invernale", "Estivo"])
    with col_s2:
        settimane_opzioni = ["Settimana 1", "Settimana 2", "Settimana 3", "Settimana 4"]
        settimana_sel = st.selectbox("Seleziona Settimana", settimane_opzioni)
    with col_s3:
        st.write("") # Spacer
        stampa_allergeni = st.checkbox("🔍 Stampa Allergeni nei File ed in Tabella", value=False)
        
    menu_sett = data["menu"].get(stagione_sel, {}).get(settimana_sel, {})
    
    if not menu_sett:
        st.warning("Nessun menù impostato per la combinazione selezionata.")
    else:
        giorni = ["Lunedì", "Martedì", "Mercoledì", "Giovedì", "Venerdì"]
        corsi = ["Primo", "Secondo", "Contorno", "Pane/Frutta"]
        
        # Costruzione Matrice per Tabella: Righe = Portate, Colonne = Giorni
        matrix_data = []
        for c in corsi:
            row = {"Portata": c}
            for g in giorni:
                piatto = menu_sett.get(g, {}).get(c, "Non impostato")
                if stampa_allergeni:
                    allg = allergeni_map.get(piatto, "Nessuno")
                    if allg and allg != "Nessuno":
                        row[g] = f"{piatto}\n[Allergeni: {allg}]"
                    else:
                        row[g] = piatto
                else:
                    row[g] = piatto
            matrix_data.append(row)
            
        df_matrix = pd.DataFrame(matrix_data)
        
        st.markdown(f"<h3 style='color:#C5A067 !important; font-family:\"Cinzel\", serif; margin-top:15px;'>📋 Menù {stagione_sel} — {settimana_sel}</h3>", unsafe_allow_html=True)
        st.caption("Struttura Matrice: Giorni in colonna | Portate in riga")
            
        render_dataframe(df_matrix)
        
        st.markdown("<h3 style='color:#C5A067 !important; font-family:\"Cinzel\", serif; margin-top:25px;'>📥 Esporta Menù</h3>", unsafe_allow_html=True)
        exp_col1, exp_col2, exp_col3 = st.columns(3)
        
        # --- EXPORT EXCEL ---
        with exp_col1:
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = f"Menu_{stagione_sel}"
            ws.sheet_view.showGridLines = True
            
            # Inserimento Logo con proporzioni reali (68x90 px)
            if os.path.exists(LOGO_PATH):
                try:
                    img_xl = OpenpyxlImage(LOGO_PATH)
                    img_xl.width = 68
                    img_xl.height = 90
                    ws.add_image(img_xl, "A1")
                except Exception:
                    pass
                    
            ws["C2"] = f"SNAILY MENSA - MENU {stagione_sel.upper()} ({settimana_sel.upper()})"
            ws["C2"].font = Font(name="Calibri", size=15, bold=True, color="024B34")
            
            # Intestazione Giorni (Riga 5)
            headers_xl = ["Portata"] + giorni
            for col_i, h_text in enumerate(headers_xl, 1):
                cell = ws.cell(row=5, column=col_i, value=h_text)
                cell.font = Font(name="Calibri", size=12, bold=True, color="FFFFFF")
                cell.fill = PatternFill(start_color="024B34", end_color="024B34", fill_type="solid")
                cell.alignment = Alignment(horizontal="center", vertical="center")
                
            # Righe Portate (Righe 6-9)
            for r_i, c_name in enumerate(corsi, 6):
                c_cell = ws.cell(row=r_i, column=1, value=c_name)
                c_cell.font = Font(name="Calibri", size=11, bold=True, color="024B34")
                c_cell.fill = PatternFill(start_color="E8F2EE", end_color="E8F2EE", fill_type="solid")
                c_cell.alignment = Alignment(horizontal="center", vertical="center")
                
                for col_i, g_name in enumerate(giorni, 2):
                    cell_val = df_matrix.loc[df_matrix["Portata"] == c_name, g_name].values[0]
                    val_cell = ws.cell(row=r_i, column=col_i, value=cell_val)
                    val_cell.font = Font(name="Calibri", size=10)
                    val_cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
                    
            ws.column_dimensions["A"].width = 16
            for col_let in ["B", "C", "D", "E", "F"]:
                ws.column_dimensions[col_let].width = 26
            ws.row_dimensions[1].height = 75
            ws.row_dimensions[5].height = 26
            for r in range(6, 10):
                ws.row_dimensions[r].height = 50
                
            out_excel = io.BytesIO()
            wb.save(out_excel)
            out_excel.seek(0)
            
            st.download_button(
                label="📊 Scarica Excel (.xlsx)",
                data=out_excel,
                file_name=f"menu_snaily_{stagione_sel.lower()}_{settimana_sel.replace(' ', '_').lower()}.xlsx",
                mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
            )
            
        # --- EXPORT WORD ---
        with exp_col2:
            doc = docx.Document()
            sec = doc.sections[0]
            sec.orientation = docx.enum.section.WD_ORIENT.LANDSCAPE
            sec.page_width = Inches(11.69)
            sec.page_height = Inches(8.27)
            sec.left_margin = Inches(0.5)
            sec.right_margin = Inches(0.5)
            
            # Intestazione con Logo ridimensionato
            p_hdr = doc.add_paragraph()
            p_hdr.alignment = WD_ALIGN_PARAGRAPH.CENTER
            if os.path.exists(LOGO_PATH):
                p_hdr.add_run().add_picture(LOGO_PATH, width=Inches(0.9))
                
            p_title = doc.add_paragraph()
            p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run_t = p_title.add_run(f"SNAILY MENSA - MENÙ {stagione_sel.upper()} ({settimana_sel.upper()})")
            run_t.bold = True
            run_t.font.size = Pt(16)
            run_t.font.color.rgb = RGBColor(2, 75, 52)
            
            # Tabella Matrice
            t_word = doc.add_table(rows=5, cols=6)
            t_word.style = 'Table Grid'
            t_word.autofit = False
            
            widths_w = [Inches(1.4), Inches(1.85), Inches(1.85), Inches(1.85), Inches(1.85), Inches(1.85)]
            
            hdr_w = t_word.rows[0].cells
            hdr_w[0].text = "Portata"
            for i, g in enumerate(giorni, 1):
                hdr_w[i].text = g
                
            from docx.oxml import parse_xml
            from docx.oxml.ns import nsdecls
            for i, cell in enumerate(hdr_w):
                cell.width = widths_w[i]
                shading = parse_xml(r'<w:shd {} w:fill="024B34"/>'.format(nsdecls('w')))
                cell._tc.get_or_add_tcPr().append(shading)
                for p in cell.paragraphs:
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    for r in p.runs:
                        r.font.bold = True
                        r.font.color.rgb = RGBColor(197, 160, 103)
                        
            for r_idx, c_name in enumerate(corsi, 1):
                row_cells = t_word.rows[r_idx].cells
                row_cells[0].text = c_name
                row_cells[0].width = widths_w[0]
                shading = parse_xml(r'<w:shd {} w:fill="E8F2EE"/>'.format(nsdecls('w')))
                row_cells[0]._tc.get_or_add_tcPr().append(shading)
                p0 = row_cells[0].paragraphs[0]
                p0.alignment = WD_ALIGN_PARAGRAPH.CENTER
                if p0.runs:
                    p0.runs[0].font.bold = True
                    p0.runs[0].font.color.rgb = RGBColor(2, 75, 52)
                    
                for c_idx, g_name in enumerate(giorni, 1):
                    val_text = df_matrix.loc[df_matrix["Portata"] == c_name, g_name].values[0]
                    row_cells[c_idx].text = val_text
                    row_cells[c_idx].width = widths_w[c_idx]
                    p_c = row_cells[c_idx].paragraphs[0]
                    p_c.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    
            out_word = io.BytesIO()
            doc.save(out_word)
            out_word.seek(0)
            
            st.download_button(
                label="📄 Scarica Word (.docx)",
                data=out_word,
                file_name=f"menu_snaily_{stagione_sel.lower()}_{settimana_sel.replace(' ', '_').lower()}.docx",
                mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document"
            )

        # --- EXPORT PDF ---
        with exp_col3:
            out_pdf = io.BytesIO()
            doc_pdf = SimpleDocTemplate(
                out_pdf,
                pagesize=landscape(A4),
                leftMargin=25,
                rightMargin=25,
                topMargin=20,
                bottomMargin=20
            )
            styles = getSampleStyleSheet()
            elements = []
            
            if os.path.exists(LOGO_PATH):
                img_pdf = ReportLabImage(LOGO_PATH, width=52, height=70)
                img_pdf.hAlign = 'CENTER'
                elements.append(img_pdf)
                elements.append(Spacer(1, 8))
                
            t_style = ParagraphStyle(
                'PdfTitle',
                parent=styles['Title'],
                fontName='Helvetica-Bold',
                fontSize=16,
                textColor=colors.HexColor('#024B34'),
                alignment=1,
                spaceAfter=12
            )
            elements.append(Paragraph(f"<b>SNAILY MENSA - MENÙ {stagione_sel.upper()} ({settimana_sel.upper()})</b>", t_style))
            
            c_style = ParagraphStyle('TableCell', fontName='Helvetica', fontSize=9, leading=11, alignment=1)
            h_style = ParagraphStyle('TableHdr', fontName='Helvetica-Bold', fontSize=10, textColor=colors.HexColor('#C5A067'), alignment=1)
            cr_style = ParagraphStyle('TableCourse', fontName='Helvetica-Bold', fontSize=10, textColor=colors.HexColor('#024B34'), alignment=1)
            
            pdf_table_data = []
            hdr_r = [Paragraph("<b>Portata</b>", h_style)] + [Paragraph(f"<b>{g}</b>", h_style) for g in giorni]
            pdf_table_data.append(hdr_r)
            
            for c_name in corsi:
                r_list = [Paragraph(f"<b>{c_name}</b>", cr_style)]
                for g_name in giorni:
                    val_str = df_matrix.loc[df_matrix["Portata"] == c_name, g_name].values[0].replace("\n", "<br/>")
                    r_list.append(Paragraph(val_str, c_style))
                pdf_table_data.append(r_list)
                
            t_pdf = Table(pdf_table_data, colWidths=[95, 138, 138, 138, 138, 138])
            t_pdf.setStyle(TableStyle([
                ('BACKGROUND', (0,0), (-1,0), colors.HexColor('#024B34')),
                ('BACKGROUND', (0,1), (0,-1), colors.HexColor('#E8F2EE')),
                ('ALIGN', (0,0), (-1,-1), 'CENTER'),
                ('VALIGN', (0,0), (-1,-1), 'MIDDLE'),
                ('GRID', (0,0), (-1,-1), 1, colors.HexColor('#C5A067')),
                ('TOPPADDING', (0,0), (-1,-1), 8),
                ('BOTTOMPADDING', (0,0), (-1,-1), 8),
            ]))
            
            elements.append(t_pdf)
            doc_pdf.build(elements)
            out_pdf.seek(0)
            
            st.download_button(
                label="🖨️ Scarica PDF (.pdf)",
                data=out_pdf,
                file_name=f"menu_snaily_{stagione_sel.lower()}_{settimana_sel.replace(' ', '_').lower()}.pdf",
                mime="application/pdf"
            )

# ==========================================
# PAGINA 2: IMPOSTA MENU SETTIMANALE
# ==========================================
elif menu_option == "✏️ Imposta Menù Settimanale":
    st.markdown("""
    <div class="info-card">
        <h2 style='color:#C5A067 !important; font-family:\"Cinzel\", serif; margin:0 0 5px 0;'>✏️ Impostazione Menù Settimanale</h2>
        <p style='margin:0;'>Seleziona i piatti per ciascun giorno e settimana per comporre il menù ufficiale.</p>
    </div>
    """, unsafe_allow_html=True)
    
    col1, col2 = st.columns(2)
    with col1:
        stagione = st.selectbox("Stagione", ["Invernale", "Estivo"])
    with col2:
        settimana = st.selectbox("Settimana", ["Settimana 1", "Settimana 2", "Settimana 3", "Settimana 4"])
        
    portate_list = data["portate"]
    primi = [p["nome"] for p in portate_list if p["categoria"] == "Primo"]
    secondi = [p["nome"] for p in portate_list if p["categoria"] == "Secondo"]
    contorni = [p["nome"] for p in portate_list if p["categoria"] == "Contorno"]
    pane_frutta = [p["nome"] for p in portate_list if p["categoria"] == "Pane/Frutta"]
    
    giorni = ["Lunedì", "Martedì", "Mercoledì", "Giovedì", "Venerdì"]
    curr_menu = data["menu"].get(stagione, {}).get(settimana, {})
    new_week_menu = {}
    
    with st.form("form_imposta_menu"):
        st.subheader(f"Composizione {settimana} - Menù {stagione}")
        for g in giorni:
            st.markdown(f"<h4 style='color:#C5A067 !important; font-family:\"Cinzel\", serif; margin-top:10px;'>📅 {g}</h4>", unsafe_allow_html=True)
            g_curr = curr_menu.get(g, {})
            
            c1, c2, c3, c4 = st.columns(4)
            with c1:
                idx_p = primi.index(g_curr.get("Primo")) if g_curr.get("Primo") in primi else 0
                sel_p = st.selectbox(f"Primo ({g})", primi, index=idx_p if primi else 0)
            with c2:
                idx_s = secondi.index(g_curr.get("Secondo")) if g_curr.get("Secondo") in secondi else 0
                sel_s = st.selectbox(f"Secondo ({g})", secondi, index=idx_s if secondi else 0)
            with c3:
                idx_c = contorni.index(g_curr.get("Contorno")) if g_curr.get("Contorno") in contorni else 0
                sel_c = st.selectbox(f"Contorno ({g})", contorni, index=idx_c if contorni else 0)
            with c4:
                idx_f = pane_frutta.index(g_curr.get("Pane/Frutta")) if g_curr.get("Pane/Frutta") in pane_frutta else 0
                sel_f = st.selectbox(f"Pane/Frutta ({g})", pane_frutta, index=idx_f if pane_frutta else 0)
                
            new_week_menu[g] = {
                "Primo": sel_p,
                "Secondo": sel_s,
                "Contorno": sel_c,
                "Pane/Frutta": sel_f
            }
            st.divider()
            
        btn_salva = st.form_submit_button("💾 Salva Menù Settimanale")
        if btn_salva:
            if stagione not in data["menu"]:
                data["menu"][stagione] = {}
            data["menu"][stagione][settimana] = new_week_menu
            save_data(data)
            st.success(f"Menù {stagione} - {settimana} salvato con successo!")

# ==========================================
# PAGINA 3: GESTISCI ANAGRAFICA PORTATE (CRUD)
# ==========================================
elif menu_option == "🥦 Gestisci Anagrafica Portate":
    st.markdown("""
    <div class="info-card">
        <h2 style='color:#C5A067 !important; font-family:\"Cinzel\", serif; margin:0 0 5px 0;'>🥦 Gestione Anagrafica Portate</h2>
        <p style='margin:0;'>Aggiungi, modifica o elimina piatti dall'Anagrafica Generale delle Scuole Snaily.</p>
    </div>
    """, unsafe_allow_html=True)
    
    df_portate = pd.DataFrame(data["portate"])
    render_dataframe(df_portate)
    
    tab_add, tab_edit, tab_del = st.tabs(["➕ Aggiungi Portata", "✏️ Modifica Portata", "❌ Elimina Portata"])
    
    # --- AGGIUNGI ---
    with tab_add:
        with st.form("form_add_portata"):
            st.subheader("Inserisci Nuova Portata")
            col1, col2 = st.columns(2)
            with col1:
                cat = st.selectbox("Categoria", ["Primo", "Secondo", "Contorno", "Pane/Frutta"])
                nome = st.text_input("Nome Piatto / Portata")
                allergeni = st.text_input("Allergeni (es. Glutine, Uova)", value="Nessuno")
            with col2:
                g_nido = st.number_input("Grammatura Nido (g)", min_value=0, value=40)
                g_inf = st.number_input("Grammatura Infanzia (g)", min_value=0, value=50)
                g_prim = st.number_input("Grammatura Primaria (g)", min_value=0, value=60)
                
            btn_add = st.form_submit_button("➕ Salva Nuova Portata")
            if btn_add:
                if not nome:
                    st.error("Inserisci il nome del piatto.")
                else:
                    new_id = max([p["id"] for p in data["portate"]], default=0) + 1
                    data["portate"].append({
                        "id": new_id,
                        "categoria": cat,
                        "nome": nome,
                        "allergeni": allergeni,
                        "nido": g_nido,
                        "infanzia": g_inf,
                        "primaria": g_prim
                    })
                    save_data(data)
                    st.success(f"Portata '{nome}' aggiunta con successo!")
                    st.rerun()

    # --- MODIFICA ---
    with tab_edit:
        st.subheader("Modifica Portata Esistente")
        portate_dict = {f"[{p['id']}] {p['nome']} ({p['categoria']})": p for p in data["portate"]}
        sel_edit = st.selectbox("Seleziona Portata da Modificare", list(portate_dict.keys()))
        
        if sel_edit:
            p_to_edit = portate_dict[sel_edit]
            with st.form("form_edit_portata"):
                c1, c2 = st.columns(2)
                with c1:
                    e_cat = st.selectbox("Categoria", ["Primo", "Secondo", "Contorno", "Pane/Frutta"], index=["Primo", "Secondo", "Contorno", "Pane/Frutta"].index(p_to_edit["categoria"]))
                    e_nome = st.text_input("Nome Piatto", value=p_to_edit["nome"])
                    e_all = st.text_input("Allergeni", value=p_to_edit["allergeni"])
                with c2:
                    e_nid = st.number_input("Grammatura Nido (g)", value=p_to_edit["nido"])
                    e_inf = st.number_input("Grammatura Infanzia (g)", value=p_to_edit["infanzia"])
                    e_pri = st.number_input("Grammatura Primaria (g)", value=p_to_edit["primaria"])
                    
                btn_save_edit = st.form_submit_button("💾 Salva Modifiche")
                if btn_save_edit:
                    for p in data["portate"]:
                        if p["id"] == p_to_edit["id"]:
                            p["categoria"] = e_cat
                            p["nome"] = e_nome
                            p["allergeni"] = e_all
                            p["nido"] = e_nid
                            p["infanzia"] = e_inf
                            p["primaria"] = e_pri
                            break
                    save_data(data)
                    st.success("Modifiche salvate con successo!")
                    st.rerun()

    # --- ELIMINA ---
    with tab_del:
        st.subheader("Elimina Portata")
        portate_dict_del = {f"[{p['id']}] {p['nome']} ({p['categoria']})": p for p in data["portate"]}
        sel_del = st.selectbox("Seleziona Portata da Eliminare", list(portate_dict_del.keys()))
        
        if sel_del:
            p_to_del = portate_dict_del[sel_del]
            if st.button("❌ Conferma Eliminazione"):
                data["portate"] = [p for p in data["portate"] if p["id"] != p_to_del["id"]]
                save_data(data)
                st.success(f"Portata '{p_to_del['nome']}' eliminata.")
                st.rerun()
